#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
批量生成 LLM 测试用例（C 组 / D 组）

用法：
    export DEEPSEEK_API_KEY=sk-xxxx
    export DASHSCOPE_API_KEY=sk-xxxx
    python docs/tools/gen_llm_tests.py --model A          # 只跑模型 A
    python docs/tools/gen_llm_tests.py --model B          # 只跑模型 B
    python docs/tools/gen_llm_tests.py --model A --only GCD BinarySearch

设计原则（不要改）：
  · 同一份 Prompt 模板，只替换【包名】和【被测类源码】
  · 单轮、不重试业务逻辑、不给原始测试
  · 模型输出原样保存，只去掉 markdown 代码围栏、只改类名与包名
  · 每次调用的参数、返回的 model 字段、token 用量全部记录到 CSV
"""

import argparse
import csv
import json
import os
import re
import sys
import time
import urllib.request
import urllib.error
from datetime import datetime

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
JAVA_PROJECT = os.path.join(REPO, "java-project")
MAIN_DIR = os.path.join(JAVA_PROJECT, "src", "main", "java")
TEST_DIR = os.path.join(JAVA_PROJECT, "src", "test", "java")
CASE_FILE = os.path.join(REPO, "docs", "01-案例清单.xlsx")
LOG_DIR = os.path.join(REPO, "docs", "raw")

# ----------------------------------------------------------------------------
# 模型配置：两款都必须走 OpenAI 兼容接口，参数才能固定并记录
# ----------------------------------------------------------------------------
MODELS = {
    "A": {
        "label": "模型A",
        "vendor": "DeepSeek",
        "base_url": "https://api.deepseek.com",
        "model": "deepseek-chat",
        "key_env": "DEEPSEEK_API_KEY",
        "temperature": 0.0,
        "top_p": 1.0,
        "max_tokens": 4096,
        "suffix": "LlmATest",
    },
    "B": {
        "label": "模型B",
        "vendor": "阿里通义千问",
        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1",
        "model": "qwen-max",
        "key_env": "DASHSCOPE_API_KEY",
        "temperature": 0.0,
        "top_p": 1.0,
        "max_tokens": 4096,
        "suffix": "LlmBTest",
    },
}

# ----------------------------------------------------------------------------
# Prompt 模板：与 docs/03-Prompt原文.md 必须逐字一致，改这里就同步改那里
# ----------------------------------------------------------------------------
PROMPT_TEMPLATE = """你是一名资深 Java 测试工程师。请为下面这个 Java 类编写 JUnit 5 单元测试。

【硬性约束】
1. 仅使用 JUnit 5（org.junit.jupiter.api），断言使用 org.junit.jupiter.api.Assertions。
   如需参数化测试，使用 @ParameterizedTest 配合 @ValueSource / @CsvSource。
2. 测试代码必须能通过 javac 编译。只能使用项目 pom.xml 中已有的依赖：
   JUnit Jupiter（org.junit.jupiter.api）、AssertJ（org.assertj.core.api）、
   Apache Commons Lang3（org.apache.commons.lang3）。
   不得引入 pom.xml 中不存在的任何新依赖。
3. 不得修改被测类；只能调用被测类中真实存在的公开方法，不得虚构方法或重载。
4. 每个测试方法只验证一个行为；方法名使用 should_期望结果_when_条件 的形式。
5. 若被测方法对非法输入抛出异常，使用 assertThrows 断言具体的异常类型。
6. 不要编写 main 方法，不要输出任何解释性文字。

【输出格式】
只输出一个 java 代码块，内容为一个完整可编译的测试类：
- 包名与下面给出的包名保持一致
- 类名为 <被测类名>Test
- 所有需要的 import 必须写全

【包名】
{pkg}

【被测类源码】
{source}"""


def load_cases():
    """从清单里读出 (包名, 类名) 列表"""
    from openpyxl import load_workbook

    ws = load_workbook(CASE_FILE)["案例清单"]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], int):
            continue
        fqcn = str(row[1]).strip()
        pkg, _, cls = fqcn.rpartition(".")
        out.append((pkg, cls, fqcn))
    return out


def read_source(pkg, cls):
    path = os.path.join(MAIN_DIR, pkg.replace(".", os.sep), cls + ".java")
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    with open(path, encoding="utf-8") as f:
        return f.read()


def call_api(cfg, prompt):
    """一次调用，返回 (文本, 元信息)"""
    key = os.environ.get(cfg["key_env"])
    if not key:
        raise SystemExit("缺少环境变量 %s" % cfg["key_env"])

    body = {
        "model": cfg["model"],
        "messages": [{"role": "user", "content": prompt}],
        "temperature": cfg["temperature"],
        "top_p": cfg["top_p"],
        "max_tokens": cfg["max_tokens"],
        "stream": False,
    }
    req = urllib.request.Request(
        cfg["base_url"].rstrip("/") + "/chat/completions",
        data=json.dumps(body).encode("utf-8"),
        headers={
            "Content-Type": "application/json",
            "Authorization": "Bearer " + key,
        },
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=180) as resp:
        data = json.loads(resp.read().decode("utf-8"))

    text = data["choices"][0]["message"]["content"]
    usage = data.get("usage") or {}
    meta = {
        "returned_model": data.get("model", ""),
        "prompt_tokens": usage.get("prompt_tokens", ""),
        "completion_tokens": usage.get("completion_tokens", ""),
    }
    return text, meta


def extract_java(text):
    """从回复里取出 java 代码块；取不到就当作整段是代码"""
    m = re.search(r"```(?:java)?\s*\n(.*?)```", text, re.S)
    code = m.group(1) if m else text
    return code.strip() + "\n"


def normalize(code, pkg, cls, suffix):
    """只做两件事：改包名、改类名。其余一律不动。"""
    # 包名：强制与目标一致
    if re.search(r"^\s*package\s+[\w.]+;", code, re.M):
        code = re.sub(r"^\s*package\s+[\w.]+;", "package %s;" % pkg, code, count=1, flags=re.M)
    else:
        code = "package %s;\n\n" % pkg + code

    # 类名：<被测类名>Test -> <被测类名><suffix>
    code = re.sub(
        r"class\s+%sTest\b" % re.escape(cls),
        "class %s%s" % (cls, suffix),
        code,
    )
    # 兜底：模型可能写成别的类名，把第一个 Test 类名统一掉
    if not re.search(r"class\s+%s%s\b" % (re.escape(cls), suffix), code):
        code = re.sub(r"class\s+(\w+Test)\b", "class %s%s" % (cls, suffix), code, count=1)
    return code


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--model", required=True, choices=["A", "B"])
    ap.add_argument("--only", nargs="*", help="只跑指定类名")
    ap.add_argument("--dry-run", action="store_true", help="只拼 Prompt，不调接口")
    args = ap.parse_args()

    cfg = MODELS[args.model]
    suffix = cfg["suffix"]
    cases = load_cases()
    if args.only:
        wanted = {s.lower() for s in args.only}
        cases = [c for c in cases if c[1].lower() in wanted]

    os.makedirs(LOG_DIR, exist_ok=True)
    log_path = os.path.join(LOG_DIR, "llm-gen-log-%s.csv" % args.model)
    new_log = not os.path.exists(log_path)
    logf = open(log_path, "a", newline="", encoding="utf-8-sig")
    writer = csv.writer(logf)
    if new_log:
        writer.writerow([
            "案例", "模型代号", "厂商标识", "请求模型名", "返回模型名",
            "temperature", "top_p", "max_tokens",
            "请求时间", "prompt_tokens", "completion_tokens", "状态",
        ])

    成功 = 失败 = 0
    for pkg, cls, fqcn in cases:
        try:
            src = read_source(pkg, cls)
        except FileNotFoundError:
            print("  ✗ 找不到源码：%s" % fqcn)
            失败 += 1
            continue

        prompt = PROMPT_TEMPLATE.format(pkg=pkg, source=src)
        if args.dry_run:
            print("  [dry] %s  (prompt %d 字符)" % (fqcn, len(prompt)))
            continue

        ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        try:
            raw, meta = call_api(cfg, prompt)
            code = normalize(extract_java(raw), pkg, cls, suffix)
            out_dir = os.path.join(TEST_DIR, pkg.replace(".", os.sep))
            os.makedirs(out_dir, exist_ok=True)
            out_path = os.path.join(out_dir, "%s%s.java" % (cls, suffix))
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(code)
            writer.writerow([
                fqcn, args.model, cfg["vendor"], cfg["model"], meta["returned_model"],
                cfg["temperature"], cfg["top_p"], cfg["max_tokens"], ts,
                meta["prompt_tokens"], meta["completion_tokens"], "ok",
            ])
            print("  ✓ %-58s -> %s" % (fqcn, os.path.basename(out_path)))
            成功 += 1
        except urllib.error.HTTPError as e:
            detail = e.read().decode("utf-8", "ignore")[:200]
            writer.writerow([fqcn, args.model, cfg["vendor"], cfg["model"], "", cfg["temperature"],
                             cfg["top_p"], cfg["max_tokens"], ts, "", "", "HTTP %s %s" % (e.code, detail)])
            print("  ✗ %-58s HTTP %s" % (fqcn, e.code))
            失败 += 1
        except Exception as e:
            writer.writerow([fqcn, args.model, cfg["vendor"], cfg["model"], "", cfg["temperature"],
                             cfg["top_p"], cfg["max_tokens"], ts, "", "", "ERR %s" % e])
            print("  ✗ %-58s %s" % (fqcn, e))
            失败 += 1

        time.sleep(1.0)   # 温柔一点，避免触发限流

    logf.close()
    print("\n模型 %s（%s / %s）完成：成功 %d，失败 %d" % (args.model, cfg["vendor"], cfg["model"], 成功, 失败))
    print("调用记录：%s" % os.path.relpath(log_path, REPO))


if __name__ == "__main__":
    main()
